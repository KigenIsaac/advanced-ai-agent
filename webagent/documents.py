from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Optional


def read_text(path: str | Path) -> str:
    return Path(path).read_text(errors="replace")


def read_csv(path: str | Path) -> dict:
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f)
        rows = list(reader)
    return {
        "headers": rows[0] if rows else [],
        "rows": rows[1:] if len(rows) > 1 else [],
        "count": max(0, len(rows) - 1),
    }


def read_xlsx(path: str | Path) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    sheets = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]
        sheets[name] = {
            "headers": rows[0] if rows else [],
            "rows": rows[1:] if len(rows) > 1 else [],
            "count": max(0, len(rows) - 1),
        }
    return sheets


def read_pdf(path: str | Path) -> dict:
    try:
        import pdfplumber
        pages_text, tables = [], []
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages):
                pages_text.append(page.extract_text() or "")
                for t in page.extract_tables() or []:
                    tables.append({"page": i, "table": t})
        return {
            "pages": len(pages_text),
            "text": "\n\n".join(pages_text),
            "tables": tables,
        }
    except ImportError:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        text = "\n\n".join((p.extract_text() or "") for p in reader.pages)
        return {"pages": len(reader.pages), "text": text, "tables": []}


def read_docx(path: str | Path) -> dict:
    from docx import Document
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    tables = []
    for t in doc.tables:
        tables.append([[cell.text for cell in row.cells] for row in t.rows])
    return {
        "text": "\n".join(paragraphs),
        "tables": tables,
    }


def read_any(path: str | Path) -> dict:
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        return {"kind": "pdf", **read_pdf(p)}
    if ext == ".docx":
        return {"kind": "docx", **read_docx(p)}
    if ext == ".xlsx" or ext == ".xlsm":
        return {"kind": "xlsx", "sheets": read_xlsx(p)}
    if ext == ".csv":
        return {"kind": "csv", **read_csv(p)}
    if ext in (".txt", ".md", ".html", ".htm", ".json", ".xml"):
        return {"kind": "text", "text": read_text(p)}
    return {"kind": "unknown", "error": f"Unsupported file type: {ext}"}